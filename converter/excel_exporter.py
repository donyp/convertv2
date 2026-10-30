from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from converter.models import ParseResult


def _money(value: Optional[Decimal]) -> Optional[float]:
    if value is None:
        return None
    return float(value)


def export_to_excel(result: ParseResult, output_path: str) -> str:
    wb = Workbook()

    # Summary sheet.
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Field", "Value"])
    ws_summary["A1"].font = Font(bold=True)
    ws_summary["B1"].font = Font(bold=True)

    metadata = result.metadata
    summary_rows = [
        ("Bank", metadata.bank_name),
        ("Bank Code", metadata.bank_code),
        ("Account Number", metadata.account_number),
        ("Account Name", metadata.account_name),
        ("Account Type", metadata.account_type),
        ("Currency", metadata.currency),
        ("Period Start", metadata.period_start),
        ("Period End", metadata.period_end),
        ("Branch", metadata.branch),
        ("Opening Balance", _money(metadata.opening_balance)),
        ("Closing Balance", _money(metadata.closing_balance)),
        ("Total Debit", _money(metadata.total_debit)),
        ("Total Credit", _money(metadata.total_credit)),
        ("Debit Record Count", metadata.debit_record_count),
        ("Credit Record Count", metadata.credit_record_count),
        ("Validation Status", result.validation_status),
        ("Transaction Count", len(result.transactions)),
    ]
    for key, value in summary_rows:
        ws_summary.append([key, value])

    if result.validation_messages:
        ws_summary.append([])
        ws_summary.append(["Validation Messages", ""])
        ws_summary["A" + str(ws_summary.max_row)].font = Font(bold=True)
        for msg in result.validation_messages:
            ws_summary.append(["", msg])

    for row in ws_summary.iter_rows(min_row=2, max_row=ws_summary.max_row, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, date):
                cell.number_format = "dd-mmm-yyyy"

    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 40

    # Transactions sheet.
    ws_txn = wb.create_sheet("Transactions")
    headers = [
        "Bank",
        "Account Number",
        "Account Name",
        "Currency",
        "Period Start",
        "Period End",
        "Transaction Date",
        "Effective Date",
        "Reference Number",
        "Description",
        "Branch/FT",
        "Debit",
        "Credit",
        "Balance",
        "Type",
        "Page",
        "Row",
        "Flags",
    ]
    ws_txn.append(headers)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_txn.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="1F2937")

    for txn in result.transactions:
        ws_txn.append([
            metadata.bank_name,
            metadata.account_number,
            metadata.account_name,
            metadata.currency,
            metadata.period_start,
            metadata.period_end,
            txn.transaction_date,
            txn.effective_date,
            txn.reference_number,
            txn.description,
            txn.branch_or_ft,
            _money(txn.debit),
            _money(txn.credit),
            _money(txn.balance),
            txn.transaction_type,
            txn.page_number,
            txn.row_number,
            "; ".join(txn.flags),
        ])

    money_cols = ["L", "M", "N"]
    date_cols = ["E", "F", "G", "H"]
    for col in money_cols:
        for row in ws_txn.iter_rows(min_row=2, min_col=ord(col) - 64, max_col=ord(col) - 64):
            for cell in row:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
    for col in date_cols:
        for row in ws_txn.iter_rows(min_row=2, min_col=ord(col) - 64, max_col=ord(col) - 64):
            for cell in row:
                cell.number_format = "dd-mmm-yyyy"

    ws_txn.freeze_panes = "A2"
    # Auto width (capped).
    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(header)
        for row in ws_txn.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 60))
        ws_txn.column_dimensions[letter].width = min(max_len + 2, 62)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
